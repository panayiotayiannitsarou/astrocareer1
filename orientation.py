from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def fmt(point):
    return f"{point.sign} {point.degree}°{point.minute:02d}′{point.second:02d}″"


def technical_source(chart):
    points = "\n".join(f"- {p.name}: {fmt(p)} | Οίκος {p.house or '—'}" for p in chart.points)
    cusps = "\n".join(f"- {i}ος Οίκος: {fmt(c)}" for i, c in enumerate(chart.cusps, 1))
    aspects = "\n".join(f"- {a.first}–{a.second} | {a.aspect} | orb {a.orb_text} | {a.weight} | {a.source}" for a in chart.aspects)
    return f"ΠΛΑΝΗΤΕΣ ΚΑΙ ΣΗΜΕΙΑ\n{points}\n\nΑΚΜΕΣ ΟΙΚΩΝ\n{cusps}\n\nΕΠΙΒΕΒΑΙΩΜΕΝΕΣ ΟΨΕΙΣ\n{aspects}"


def master_instruction():
    """Return the complete approved binding instruction bundled with the app."""
    path = Path(__file__).with_name("master_instruction.docx")
    document = Document(path)
    return "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())


def service_template(service):
    filename = (
        "template_child_teen.docx"
        if service == "Παιδί/έφηβος"
        else "template_adult.docx"
    )
    document = Document(Path(__file__).with_name(filename))
    return "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())


def build_prompt(chart, name, service, presentation, language, cyprus_school=False):
    language_rule = (
        "Γράψε και τα δύο παραδοτέα εξ ολοκλήρου στα Ελληνικά."
        if language == "Ελληνικά"
        else "Write both deliverables entirely in natural, clear English. Translate every title, heading, label and table field into English."
    )
    cyprus_value = "Ναι" if service == "Παιδί/έφηβος" and cyprus_school else "Όχι"
    return f"""ΕΝΤΟΛΗ ΕΚΤΕΛΕΣΗΣ ΓΙΑ ΤΗ ΣΥΓΚΕΚΡΙΜΕΝΗ ΠΕΡΙΠΤΩΣΗ

Ακολούθησε ολόκληρη τη ΔΕΣΜΕΥΤΙΚΗ ΕΝΤΟΛΗ που ακολουθεί. Οι ειδικοί υπερισχύοντες Κανόνες 0Γ και 0Δ εφαρμόζονται πριν από τη γενική δομή του Κανόνα 16.

Παράδωσε ακριβώς δύο χωριστά αρχεία Word:
1. Καθαρό Παραδοτέο πελάτη, σύντομο και απλό, με ακριβώς τη δομή του κατάλληλου ανώνυμου προτύπου που ενσωματώνεται παρακάτω και τη λιτή μορφοποίηση του εγκεκριμένου παραδείγματος GAVRIELA. Να έχει λευκό φόντο, μαύρο κείμενο, καθαρούς τίτλους και χωρίς πίνακες, τεχνική αστρολογική ορολογία, χρωματιστά πλαίσια ή σκιάσεις.
2. Εσωτερικό Τεχνικό Δελτίο Ελέγχου, αναλυτικό, με τους πίνακες τεκμηρίωσης, τα ακριβή τεχνικά δεδομένα, τον έλεγχο των στενότερων όψεων, τη δήλωση συμμόρφωσης και τον αυτοέλεγχο του καθαρού παραδοτέου.

Μην συγχωνεύσεις τα δύο αρχεία. Μην εμφανίσεις το εσωτερικό τεχνικό δελτίο στον πελάτη.

Όνομα: {name}
Τύπος υπηρεσίας: {service}
Τρόπος παρουσίασης: Απλή και πρακτική
Γλώσσα: {language}
Φοίτηση στο κυπριακό εκπαιδευτικό σύστημα: {cyprus_value}

{language_rule}

ΥΠΟΧΡΕΩΤΙΚΟ ΠΡΟΤΥΠΟ ΚΑΘΑΡΟΥ ΠΑΡΑΔΟΤΕΟΥ ΓΙΑ ΤΗΝ ΕΠΙΛΕΓΜΕΝΗ ΥΠΗΡΕΣΙΑ
Το ακόλουθο είναι πρότυπο δομής και ύφους, όχι πηγή προσωπικών ή τεχνικών δεδομένων. Αντικατάστησε τα υποδείγματα και τις αγκύλες με το πραγματικό περιεχόμενο της παρούσας περίπτωσης.
{service_template(service)}

ΠΛΗΡΗΣ ΔΕΣΜΕΥΤΙΚΗ ΕΝΤΟΛΗ
{master_instruction()}

ΕΛΕΓΜΕΝΑ ΤΕΧΝΙΚΑ ΔΕΔΟΜΕΝΑ
{technical_source(chart)}
"""


TECHNICAL_TERMS = re.compile(r"\b(?:Ήλιος|Σελήνη|Ερμής|Αφροδίτη|Άρης|Δίας|Κρόνος|Ουρανός|Ποσειδώνας|Πλούτωνας|Χείρωνας|Ωροσκόπος|Μεσουράνημα|orb|Ο[ίι]κος|σύνοδος|τρίγωνο|εξάγωνο|τετράγωνο|αντίθεση)\b|\d{1,2}°\d{1,2}[′']", re.I)


def validate_result(text, service, presentation, language, audit_text=""):
    errors = []
    low = text.lower()
    if len(text.strip()) < 700:
        errors.append("Το κείμενο είναι υπερβολικά σύντομο.")
    greek_terms = ["προφίλ", "ταλέντ", "επαγγελματικ", "τελική σύνθεση"]
    english_terms = ["profile", "talent", "career", "final synthesis"]
    required = greek_terms if language == "Ελληνικά" else english_terms
    for term in required:
        if term not in low:
            errors.append(f"Δεν εντοπίστηκε η βασική ενότητα: {term}.")
    if service != "Παιδί/έφηβος" and re.search(r"8\s*[–-]\s*12\s+(?:εβδομάδ|weeks)|εργασιακ(?:ά|ό)\s+περιβάλλον|work\s+environment", text, re.I):
        errors.append("Η έκδοση ενηλίκου περιέχει ενότητα που έχει αφαιρεθεί από τις προδιαγραφές.")
    if presentation == "Απλή και πρακτική":
        if TECHNICAL_TERMS.search(text):
            errors.append("Η απλή έκδοση περιέχει τεχνική αστρολογική ορολογία.")
        if not audit_text.strip():
            errors.append("Λείπει το εσωτερικό τεχνικό δελτίο.")
        elif not re.search(
            r"(?:Παράρτημα\s+Ελέγχου\s+Τεκμηρίωσης|"
            r"Εσωτερικό\s+Τεχνικό\s+Δελτίο(?:\s+Ελέγχου)?|"
            r"Evidence\s+Audit\s+Appendix|Internal\s+Technical\s+(?:Audit\s+)?Record)",
            audit_text,
            re.I,
        ):
            errors.append("Το τεχνικό δελτίο δεν έχει αναγνωρίσιμη ενότητα ελέγχου.")
    return errors


def validate_docx_format(client_data, audit_data):
    """Check the two-file split and the approved GAVRIELA document shape."""
    errors = []
    client = Document(BytesIO(client_data))
    audit = Document(BytesIO(audit_data))
    if client.tables:
        errors.append("Το καθαρό παραδοτέο πρέπει να έχει τη λιτή μορφή GAVRIELA χωρίς πίνακες.")
    if len(audit.tables) < 3:
        errors.append("Το εσωτερικό τεχνικό δελτίο δεν περιέχει τους απαιτούμενους πίνακες ελέγχου.")
    return errors


def docx_text(data):
    document = Document(BytesIO(data))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def _shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def prompt_docx(name, service, prompt):
    document = Document()
    sec = document.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.8)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(service).font.color.rgb = RGBColor.from_string("76538C")
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(name).bold = True
    for line in prompt.splitlines():
        if line.isupper() and len(line) < 100:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    bio = BytesIO()
    document.save(bio)
    return bio.getvalue()
